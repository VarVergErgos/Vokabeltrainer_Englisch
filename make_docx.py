import sys; sys.path.insert(0, "app")
import helptext as H
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(12)
for n, sz in (("Heading 1", 16), ("Title", 24)):
    d.styles[n].font.name = "Calibri"; d.styles[n].font.size = Pt(sz); d.styles[n].font.color.rgb = RGBColor(0x0B, 0x7A, 0x75)
d.add_paragraph(H.TITEL, style="Title")
d.add_paragraph("Stand: 11.10.2026 · Stufe 3, Sprachduschen, EXE-Bau, Suche, Installer · Alles läuft offline, die Daten liegen im Programmordner.")
d.add_heading("Inhalt", 1)
for i, (t, _) in enumerate(H.KAPITEL, 1):
    d.add_paragraph("%d. %s" % (i, t))
d.add_page_break()
for i, (t, items) in enumerate(H.KAPITEL, 1):
    d.add_heading("%d. %s" % (i, t), 1)
    for it in items:
        if isinstance(it, str):
            d.add_paragraph(it)
        elif it[0] == "liste":
            for x in it[1]:
                d.add_paragraph(x, style="List Bullet")
        else:
            p = d.add_paragraph(); r = p.add_run("Tipp: "); r.bold = True; p.add_run(it[1])
            pPr = p._p.get_or_add_pPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "FFF4CC"); pPr.append(sh)
sec = d.sections[0]; fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Seite "); 
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    run = fp.add_run()
    if tag:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); run._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)
d.save(sys.argv[1])
