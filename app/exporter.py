# Export eines Kapitels: TXT, Excel, PDF (Qt, kein Zusatzpaket) und MP3/WAV (Vorlesen).
import os, re, wave
import db
from paths import ensure

def _safe(s):
    return re.sub(r'[^A-Za-z0-9äöüÄÖÜß _-]+', "", s).strip().replace(" ", "_")[:40] or "Kapitel"

def default_name(k, ext):
    return "Kapitel_%02d_%s.%s" % (k["nr"], _safe(k["titel"]), ext)

def export_txt(conn, k, path):
    """UTF-8, Englisch und darunter Deutsch - direkt fuer MP3-Programme nutzbar."""
    z = ["Kapitel %d: %s" % (k["nr"], k["titel"]), ""]
    for e in db.entries(conn, k["id"]):
        z += [e["englisch"], e["deutsch"]]
        if e["aussprache"]:
            z.append("Aussprache: " + e["aussprache"])
        if e["beispiel_en"]:
            z += [e["beispiel_en"], e["beispiel_de"]]
        z.append("")
    with open(path, "w", encoding="utf-8", newline="\r\n") as f:
        f.write("\n".join(z))
    return path

def export_xlsx(conn, k, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Kapitel %d" % k["nr"]
    ws.append(["Englisch", "Deutsch", "Aussprache", "Beispiel Englisch", "Beispiel Deutsch"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for e in db.entries(conn, k["id"]):
        ws.append([e["englisch"], e["deutsch"], e["aussprache"], e["beispiel_en"], e["beispiel_de"]])
    for col, w in zip("ABCDE", (34, 34, 30, 44, 44)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path); return path

def _esc(s):
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def export_pdf(conn, k, path):
    """Druckfreundliche PDF-Seite mit grosser Schrift (Qt-PDF, braucht keine Zusatzprogramme)."""
    from PySide6.QtGui import QTextDocument, QPdfWriter, QPageSize, QPageLayout
    from PySide6.QtCore import QMarginsF, QUrl
    h = ["<h1>Kapitel %d: %s</h1>" % (k["nr"], _esc(k["titel"]))]
    if k["hinweis"]:
        h.append("<p><i>%s</i></p>" % _esc(k["hinweis"]))
    h.append('<table width="100%" cellspacing="0" cellpadding="6" border="1" style="border-color:#bbb">')
    h.append("<tr><th align=left>Englisch</th><th align=left>Deutsch</th></tr>")
    for e in db.entries(conn, k["id"]):
        en = "<b>%s</b>" % _esc(e["englisch"])
        if e["aussprache"]:
            en += "<br><span style='color:#555'>[%s]</span>" % _esc(e["aussprache"])
        if e["beispiel_en"]:
            en += "<br><i>%s</i>" % _esc(e["beispiel_en"])
        de = _esc(e["deutsch"])
        if e["beispiel_de"]:
            de += "<br><i>%s</i>" % _esc(e["beispiel_de"])
        h.append("<tr><td>%s</td><td>%s</td></tr>" % (en, de))
    h.append("</table>")
    for b in db.images(conn, k["id"]):
        f = os.path.join(ensure("bilder"), b["dateiname"])
        if os.path.exists(f):
            h.append('<p><img src="%s" width="560"><br>%s</p>' % (QUrl.fromLocalFile(f).toString(), _esc(b["beschreibung"])))
    doc = QTextDocument(); doc.setDefaultStyleSheet("body{font-family:'Segoe UI',Arial;font-size:13pt} h1{font-size:20pt}")
    doc.setHtml("<body>%s</body>" % "".join(h))
    w = QPdfWriter(path); w.setPageSize(QPageSize(QPageSize.A4)); w.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    doc.print_(w); return path

def export_audio(conn, k, path, speaker, voice, speed, hoer, progress=None):
    """Liest Englisch (Wort/Satz, danach Beispielsatz) vor. .mp3 (lameenc) oder .wav, mit kurzen Pausen."""
    import numpy as np, speech
    es = db.entries(conn, k["id"]); parts = []; sil = np.zeros(int(speech.SR * 0.9), dtype=np.int16)
    for i, e in enumerate(es):
        for t in (e["englisch"], e["beispiel_en"]):
            if t:
                with wave.open(speaker.wav_for(t, voice, speed, hoer)) as w:
                    parts.append(np.frombuffer(w.readframes(w.getnframes()), dtype=np.int16)); parts.append(sil)
        if progress:
            progress(i + 1, len(es))
    pcm = np.concatenate(parts) if parts else np.zeros(1, dtype=np.int16)
    if path.lower().endswith(".mp3"):
        import lameenc
        enc = lameenc.Encoder(); enc.set_bit_rate(128); enc.set_in_sample_rate(speech.SR); enc.set_channels(1); enc.set_quality(2)
        data = enc.encode(pcm.tobytes()) + enc.flush()
        with open(path, "wb") as f:
            f.write(data)
    else:
        w = wave.open(path, "wb"); w.setnchannels(1); w.setsampwidth(2); w.setframerate(speech.SR); w.writeframes(pcm.tobytes()); w.close()
    return path


# ---- MP3-Export mit Uebersetzung (Menue "MP3 erstellen"): je Kapitel eine Excel-Datei und eine MP3 im Ordner export_MP3
BAD = '<>:"/\\|?*'

def mp3_basename(k):
    """0001 + Leerzeichen + Kapitelname (hoechstens 25 Zeichen), ohne Zeichen, die Windows in Dateinamen verbietet."""
    t = "".join(c for c in (k["titel"] or "") if c not in BAD and ord(c) >= 32)
    t = re.sub(r"\s+", " ", t).strip()[:25].strip(" .")
    return ("%04d %s" % (k["nr"], t)).strip()

def mp3_rows(conn, k, beispiele=True):
    z = []
    for e in db.entries(conn, k["id"]):
        if e["englisch"] and e["deutsch"]:
            z.append((e["deutsch"], e["englisch"]))
        if beispiele and e["beispiel_en"] and e["beispiel_de"]:
            z.append((e["beispiel_de"], e["beispiel_en"]))
    return z

def mp3_write_xlsx(path, k, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook(); ws = wb.active; ws.title = ("Kapitel %d" % k["nr"])[:31]
    ws.append(["Nr", "Deutsch", "Englisch"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, (de, en) in enumerate(rows, 1):
        ws.append([i, de, en])
    ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 58; ws.column_dimensions["C"].width = 58
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)

def mp3_read_xlsx(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).worksheets[0]
    z = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        de, en = ("" if v is None else str(v).strip() for v in (r[1], r[2]))
        if de and en:
            z.append((de, en))
    return z

def _pcm(speaker, text, voice, speed, hoer):
    import numpy as np
    with wave.open(speaker.wav_for(text, voice, speed, hoer)) as w:
        return np.frombuffer(w.readframes(w.getnframes()), dtype=np.int16)

def export_mp3_paare(conn, k, ordner, speaker, voice_en, voice_de, speed, hoer, deutsch_zuerst=True,
                     pause_sprache=1.5, pause_paar=3.0, vorhandene_excel=False, beispiele=True, progress=None):
    """Legt ordner\\0001 Name.xlsx (Deutsch | Englisch) und ordner\\0001 Name.mp3 an.
    Die MP3 wird aus der Excel-Datei erzeugt (bearbeitete Excel: vorhandene_excel=True).
    Pro Zeile: erste Sprache, Pause (pause_sprache), zweite Sprache, Pause (pause_paar). Rueckgabe (mp3, xlsx, anzahl)."""
    import numpy as np, speech
    os.makedirs(ordner, exist_ok=True)
    base = os.path.join(ordner, mp3_basename(k))
    xl, mp = base + ".xlsx", base + ".mp3"
    if vorhandene_excel and os.path.exists(xl):
        rows = mp3_read_xlsx(xl)
    else:
        rows = mp3_rows(conn, k, beispiele)
        mp3_write_xlsx(xl, k, rows)
        rows = mp3_read_xlsx(xl)
    if not rows:
        raise ValueError("Kapitel %d hat keine Einträge mit Englisch und Deutsch." % k["nr"])
    p1 = np.zeros(int(speech.SR * pause_sprache), dtype=np.int16)
    p2 = np.zeros(int(speech.SR * pause_paar), dtype=np.int16)
    teile = []
    for i, (de, en) in enumerate(rows):
        a = (de, voice_de) if deutsch_zuerst else (en, voice_en)
        b = (en, voice_en) if deutsch_zuerst else (de, voice_de)
        teile += [_pcm(speaker, a[0], a[1], speed, hoer), p1, _pcm(speaker, b[0], b[1], speed, hoer), p2]
        if progress:
            progress(i + 1, len(rows))
    import lameenc
    enc = lameenc.Encoder(); enc.set_bit_rate(128); enc.set_in_sample_rate(speech.SR); enc.set_channels(1); enc.set_quality(2)
    data = enc.encode(np.concatenate(teile).tobytes()) + enc.flush()
    with open(mp, "wb") as f:
        f.write(data)
    return mp, xl, len(rows)
